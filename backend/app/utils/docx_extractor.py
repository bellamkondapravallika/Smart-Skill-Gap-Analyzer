import io
from docx import Document

def extract_text_from_docx_bytes(docx_bytes: bytes) -> str:
    """
    Extract readable text from DOCX bytes using python-docx.
    Processes paragraphs, headings, bullet lists, and tables.
    """
    try:
        doc = Document(io.BytesIO(docx_bytes))
        paragraphs = []
        
        # Read paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)
        
        # Also read any table cells
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        
        full_text = "\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("The uploaded DOCX document contains no readable text content.")
        return full_text
    except Exception as e:
        raise ValueError(f"Could not extract text from DOCX: {str(e)}")
